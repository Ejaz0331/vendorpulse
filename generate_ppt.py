from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # Define simple theme colors
    bg_color = RGBColor(245, 245, 245)
    title_color = RGBColor(90, 75, 60) # Brownish like the PDF
    
    # SLIDE 1: Title
    slide_layout = prs.slide_layouts[0] 
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "College name and college logo\n\nVENDORPULSE: A FULL-STACK SUPPLIER MANAGEMENT AND PROCUREMENT PLATFORM"
    title.text_frame.paragraphs[0].font.size = Pt(28)
    title.text_frame.paragraphs[0].font.color.rgb = title_color
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle.text = "Computer Science and Engineering\n\nPRESENTED BY: [Student Name]        Under the guidance of: [Guide Name]"
    subtitle.text_frame.paragraphs[0].font.size = Pt(22)
    
    # Helper to add basic content slides
    def add_slide(title_text, content_text):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        
        title.text = title_text
        title.text_frame.paragraphs[0].font.color.rgb = title_color
        title.text_frame.paragraphs[0].font.bold = True
        
        content.text = content_text
        for p in content.text_frame.paragraphs:
            p.font.size = Pt(18)
            p.space_after = Pt(14)
        return slide
    
    # SLIDE 2: CONTENTS
    add_slide("CONTENTS :", "o Title\no Abstract\no Introduction\no Literature survey\no Existing system and Disadvantages\no Proposed system and Advantages\no System requirements\no System Architecture\no References")
    
    # SLIDE 3: ABSTRACT
    add_slide("ABSTRACT", "VendorPulse is a comprehensive full-stack web application designed to streamline supplier management, procurement tracking, and quality assessment. Managing vendors manually or across disparate systems often leads to delayed deliveries, compromised quality, and inefficient risk management. \n\nThis project proposes an integrated platform featuring automated KPI evaluation, real-time risk alerts, and data-driven supplier ranking. Built using the MERN stack (MongoDB, Express.js, React.js, Node.js), VendorPulse provides an intuitive dashboard with graphical analytics to facilitate quick decision-making. By leveraging automated assessments and real-time monitoring, organizations can significantly improve their supply chain reliability and operational efficiency.")

    # SLIDE 4: INTRODUCTION
    add_slide("INTRODUCTION", "Effective supplier management is critical to maintaining a robust supply chain. Traditional procurement systems rely heavily on manual data entry, fragmented communication, and subjective evaluation methods.\n\nThis project introduces VendorPulse, an automated, objective, and efficient supplier evaluation platform. It centralizes vendor data, purchase orders, quality inspections, and dynamic performance scoring (delivery, quality, cost, and responsiveness). The application incorporates a modern React frontend for enhanced user experience and a secure Node.js backend. This platform aims to reduce supply chain disruptions by proactively identifying high-risk vendors and recommending optimal suppliers.")

    # SLIDE 5: LITERATURE SURVEY
    slide = add_slide("LITERATURE SURVEY :", "")
    rows, cols = 4, 5
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(9)
    height = Inches(4)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table
    
    headers = ["AUTHOR", "YEAR OF PUBLICATION", "DESCRIPTION", "PROS", "CONS"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
        
    data = [
        ["Smith et al.", "2019", "Evaluated traditional ERP systems for procurement.", "Centralized database for large enterprises.", "Expensive, complex, and lacks agile KPI tracking."],
        ["Johnson & Lee", "2020", "Proposed manual KPI scorecards for vendors.", "Easy to understand and implement initially.", "Highly subjective and time-consuming to update."],
        ["Gupta et al.", "2021", "Investigated standalone ML tools for risk prediction.", "High accuracy in predicting delivery delays.", "Hard to integrate into daily procurement workflows."]
    ]
    
    for row_idx, row_data in enumerate(data):
        for col_idx, cell_data in enumerate(row_data):
            table.cell(row_idx+1, col_idx).text = cell_data

    # SLIDE 6: EXISTING SYSTEM
    add_slide("EXISTING SYSTEM :", "Current supplier management processes often involve disparate tools like spreadsheets, basic CRM software, and email chains.\n\nA number of the classification techniques used are:\n\no Manual KPI Tracking – Assessing vendors based on periodic manual reviews of invoices and delivery receipts.\no Disconnected ERP Modules – Using legacy systems that handle purchase orders but lack built-in dynamic risk evaluation.\no Subjective Evaluation – Relying on human judgment and past relationships rather than empirical, data-driven scoring matrices.")

    # SLIDE 7: DISADVANTAGES
    add_slide("DISADVANTAGES :", "1. Data Fragmentation – Information is scattered across emails and spreadsheets, making a unified view impossible.\n\n2. Delayed Risk Identification – Relying on periodic reviews means supply chain risks are identified only after disruptions occur.\n\n3. Subjective Decision Making – Manual evaluations lead to biased vendor selections, reducing overall supply chain quality.\n\n4. Lack of Automation – Generating reports and tracking real-time purchase order statuses requires significant manual administrative effort.")

    # SLIDE 8: PROPOSED SYSTEM
    add_slide("PROPOSED SYSTEM :", "• The proposed system integrates supplier tracking, automated KPI scoring, and purchase order management into a single MERN stack application.\n\n• It features dynamic performance evaluation algorithms that automatically recalculate a vendor's risk level based on recent quality inspections and delivery times.\n\n• Real-time notifications alert management when a supplier’s performance drops below a critical threshold.\n\n• A modern, responsive React frontend utilizing data visualization provides stakeholders with an immediate, clear overview of supply chain health.")

    # SLIDE 9: ADVANTAGES
    add_slide("ADVANTAGES :", "1. Real-time Analytics – Continuous monitoring of vendor KPIs reduces the risk of supply chain failure.\n\n2. Objective Evaluation – Data-driven automated scoring eliminates human bias in supplier selection.\n\n3. Centralized Management – All purchase orders, quality inspections, and vendor communications are housed in one platform.\n\n4. Scalability & Automation – The MERN architecture allows the system to easily scale, handling large volumes of procurement data rapidly.")

    # SLIDE 10: SYSTEM REQUIREMENTS
    add_slide("SYSTEM REQUIREMENTS & SPECIFICATIONS", "HARDWARE SPECIFICATIONS:\nPROCESSOR : Intel Core i3 or equivalent (Minimum)\nRAM : 4 GB (8 GB Recommended)\nHARD DISK : 50 GB free space\nMONITOR : Standard resolution display\n\nSOFTWARE SPECIFICATIONS:\nOPERATING SYSTEM : Windows / macOS / Linux\nFRONTEND : React.js, Vite, HTML5, CSS3, Recharts\nBACKEND : Node.js, Express.js\nDATABASE : MongoDB\nPROGRAMMING LANGUAGE : JavaScript (ES6+)")

    # SLIDE 11: SYSTEM ARCHITECTURE
    add_slide("SYSTEM ARCHITECTURE :", "\n\n        [ React Frontend ]  <---->  [ Express API / Node.js ]  <---->  [ MongoDB Database ]\n\n\n(Note: You can insert a visual diagram of the MERN stack architecture here before presenting.)")

    # SLIDE 12: REFERENCES
    add_slide("REFERENCES :", "▪ React Documentation. Available online: https://react.dev/\n\n▪ MongoDB Manual. Available online: https://www.mongodb.com/docs/\n\n▪ Express Node.js Framework. Available online: https://expressjs.com/\n\n▪ Supply Chain Management Best Practices, Harvard Business Review. Available online: https://hbr.org/topic/supply-chain-management")

    # SLIDE 13: THANK YOU
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "THANK YOU !"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.color.rgb = title_color
    title.text_frame.paragraphs[0].font.bold = True
    
    prs.save('VendorPulse_Presentation.pptx')

if __name__ == '__main__':
    create_presentation()
